# Membuat file docx tiruan template soal OSCE (Lampiran III pedoman) untuk menguji fitur impor rubrik.
import os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'docx')
PER = ["Menunjukkan {} perilaku profesional meliputi:", "1) Mengucapkan salam pembuka kepada pasien", "2) Bersikap ramah", "3) Mempertahankan kontak mata", "4) Menggunakan bahasa sederhana dan santun", "5) Melakukan tindakan konseling dengan tanggung jawab", "6) Mengucapkan salam penutup"]

def isi_sel(cell, teks, strike_kecuali=None):
    cell.text = ''
    baris = teks if isinstance(teks, list) else [teks]
    p = cell.paragraphs[0]
    for i, b in enumerate(baris):
        if i: p = cell.add_paragraph()
        if strike_kecuali is not None:
            # "Stasion 1 / 2 / 3 ..." dengan angka selain strike_kecuali dicoret
            p.add_run('Stasion ')
            for n in range(1, b + 1):
                r = p.add_run(f'{n}' + (' / ' if n < b else ''))
                if n != strike_kecuali: r.font.strike = True
                else: r.bold = True
            p.add_run(' (stasion istirahat)')
        else:
            p.add_run(b)

def kuning(row):
    for c in row.cells:
        tcPr = c._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'FFFF00'); tcPr.append(shd)

def buat(nama, no, jml_st, judul, butir, judul_heading=True, nomor_coret=True, sub_header=None):
    d = Document()
    if judul_heading: d.add_heading(f'Stasion – {no}.', 1)
    t = d.add_table(rows=4, cols=3); t.style = 'Table Grid'
    lab = [('1.', 'Nomor station'), ('2.', 'Judul stasion'), ('3.', 'Waktu yang dibutuhkan'), ('7.', 'Instruksi peserta ujian')]
    for i, (a, b) in enumerate(lab):
        t.rows[i].cells[0].text = a; t.rows[i].cells[1].text = b
    if nomor_coret: isi_sel(t.rows[0].cells[2], jml_st, strike_kecuali=no)
    else: t.rows[0].cells[2].text = 'Stasion ' + ' / '.join(str(i) for i in range(1, jml_st + 1))
    t.rows[1].cells[2].text = judul; t.rows[2].cells[2].text = '10 menit'
    t.rows[3].cells[2].text = 'Skenario klinik: RAHASIA – tidak boleh ikut terimpor.'
    d.add_paragraph('RUBRIK PENILAIAN')
    r = d.add_table(rows=1, cols=8); r.style = 'Table Grid'
    for i, h in enumerate(['NO', 'Kompetensi *) (diantara Kompetensi yang diujikan)', 'Skor = 0, jika', 'Skor = 1, jika', 'Skor = 2, jika', 'Skor = 3, jika', 'BOBOT NILAI (1-3)', 'SKOR MAKS (skor tertinggi x bobot)']):
        r.rows[0].cells[i].text = h
    for b in butir:
        row = r.add_row()
        vals = [b.get('no', ''), b['k'], b.get('s0', ''), b.get('s1', ''), b.get('s2', ''), b.get('s3', ''), b.get('bobot', ''), '']
        for i, v in enumerate(vals): isi_sel(row.cells[i], v)
        if b.get('alt'): kuning(row)
    d.add_paragraph('Keterangan : Bobot 1 (Mudah), Bobot 2 (Sedang), Bobot 3 (Sulit)')
    g = d.add_table(rows=2, cols=4); g.style = 'Table Grid'
    for i, h in enumerate(['TIDAK LULUS', 'BORDERLINE', 'LULUS', 'SUPERIOR']): g.rows[0].cells[i].text = h
    d.save(os.path.join(OUT, nama)); print('ditulis', nama)

perilaku = dict(k='Mampu menunjukkan perilaku profesional', s0='Tidak memiliki perilaku profesional sama sekali',
                s1=[PER[0].format('1-2')] + PER[1:], s2=[PER[0].format('3-4')] + PER[1:], s3=[PER[0].format('5-6')] + PER[1:], bobot='2')
# D3 St 2 Gastritis (lengkap, heading + coret)
buat('Soal_D3_St2_Gastritis.docx', 2, 8, 'Sistem Pencernaan: Gastritis', [
    dict(no='1', k='Mampu melakukan intervensi pemberian makanan', s0='Tidak menentukan intervensi pemberian makanan yang tepat',
         s1=['Menentukan intervensi 1 dari 4 aspek:', '1) Jenis diet', '2) Bentuk makanan', '3) Rute', '4) Prinsip diet'],
         s2=['Menentukan intervensi 2-3 dari 4 aspek:', '1) Jenis diet', '2) Bentuk makanan', '3) Rute', '4) Prinsip diet'],
         s3=['Menentukan intervensi semua (4) aspek:', '1) Jenis diet', '2) Bentuk makanan', '3) Rute', '4) Prinsip diet'], bobot='3'),
    dict(no='2', k='Melakukan edukasi gizi terkait kasus Gastritis (versi revisi 2027)', s0='Tidak melakukan edukasi gizi',
         s1='Melakukan 1-2 dari 5 langkah', s2='Melakukan 3-4 dari 5 langkah', s3='Melakukan 5 langkah', bobot='3'),
    dict(no='3', **perilaku)])
# D3 St 3: ada baris alternatif kuning, tanpa heading -> nomor dari sel "Nomor station" yang dicoret
buat('Soal_Dislipidemia.docx', 3, 8, 'Kardiovaskuler: Dislipidemia', [
    dict(no='1', k='Mampu melakukan pengukuran antropometri (BB dan TB) dan data objektif', s0='Tidak melakukan', s1='1 pengukuran benar', s2='2 pengukuran, 1 benar', s3='2 pengukuran benar', bobot='3'),
    dict(no='', k='(atau hanya 1 pengukuran antropometri, misal tinggi badan)', s0='Tidak mengukur TB', s1='1-2 langkah', s2='3 langkah', s3='4 langkah', bobot='3', alt=True),
    dict(no='2', k='Mampu menetapkan diagnosis gizi', s0='Tidak menetapkan', s1='1 aspek PES', s2='2 aspek PES', s3='3 aspek PES', bobot='1'),
    dict(no='3', **perilaku)], judul_heading=False)
# D3 tanpa penanda nomor sama sekali -> harus dipilih manual / dari nama file
buat('St5 Balita Stunting.docx', 5, 8, 'Anak: Balita stunting', [
    dict(no='1', k='Mampu melakukan pengukuran antropometri', s0='Tidak', s1='a', s2='b', s3='c', bobot='3'),
    dict(no='2', k='Melakukan edukasi gizi balita stunting', s0='Tidak', s1='1-2 langkah', s2='3-4 langkah', s3='6 langkah', bobot=''),
    dict(no='3', **perilaku)], judul_heading=False, nomor_coret=False)
# S.Tr St 8 Katering dengan sub-butir 1.1-1.4
buat('Soal_STR_St8_Katering.docx', 8, 9, 'Penyelenggaraan Makanan: Katering', [
    dict(no='1', k='Mampu melakukan Pendokumentasian :'),
    dict(no='1.1.', k='Perencanaan kebutuhan bahan makanan', s0='Tidak menghitung', s1='konsumen + porsi', s2='+ frekuensi', s3='+ putaran siklus', bobot='2'),
    dict(no='1.2', k='Kegiatan sebagai penyelia', s0='Tidak', s1='tanpa instrumen', s2='instrumen sebagian', s3='instrumen lengkap', bobot='2'),
    dict(no='1.3', k='Uji mutu fisik dan organoleptik', s0='Tidak', s1='salah satu', s2='keduanya tanpa SOP', s3='keduanya dengan SOP', bobot='3'),
    dict(no='1.4', k='Kesimpulan hasil praktik', s0='Tidak', s1='1 dari 3', s2='2 dari 3', s3='3 dari 3', bobot='1'),
    dict(no='2', **perilaku)])
